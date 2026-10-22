"""Monitoramento: lê a lista exportada de matrículas e acha os colaboradores.

A plataforma exporta a lista em CSV, Excel (.xlsx) ou numa tabela HTML salva
como .xls; as três formas são lidas aqui. Os nomes são comparados sem acento,
sem maiúsculas e sem espaços repetidos.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from .gerador import PARTICULAS, sem_acento


@dataclass
class Matricula:
    aluno: str
    curso: str
    progresso: float | None  # 0 a 100; None se a lista não trouxer
    texto_progresso: str = ""
    email: str = ""

    @property
    def concluida(self) -> bool:
        return self.progresso is not None and self.progresso >= 99.995


@dataclass
class Resultado:
    procurado: str
    matriculas: list[Matricula]


def normalizar(nome: str) -> str:
    return re.sub(r"\s+", " ", sem_acento(nome or "")).strip().lower()


def nomes_colados(texto: str) -> list[str]:
    """Um nome por linha (também aceita ; e tabulação), sem repetidos e sem números soltos."""
    nomes, vistos = [], set()
    for parte in re.split(r"[\n;\t]+", texto or ""):
        nome = re.sub(r"^\s*\d+\s*[-.)]?\s*", "", parte).strip(" ,")  # "1. Maria" / "1 - Maria"
        chave = normalizar(nome)
        if chave and chave not in vistos:
            vistos.add(chave)
            nomes.append(re.sub(r"\s+", " ", nome))
    return nomes


def ler_progresso(texto: str) -> float | None:
    numero = re.search(r"\d+(?:[.,]\d+)?", str(texto or ""))
    if not numero:
        return None
    valor = float(numero.group().replace(",", "."))
    if "%" not in str(texto) and valor <= 1 and "." in numero.group():
        valor *= 100  # Excel com o número em fração (0,5 = 50%)
    return valor


# ---------- leitura do arquivo exportado ----------
class _TabelaHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.linhas: list[list[str]] = []
        self._linha: list[str] | None = None
        self._celula: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._celula is not None and self._linha is not None:
            self._linha.append(" ".join("".join(self._celula).split()))
            self._celula = None
        elif tag == "tr" and self._linha is not None:
            if any(self._linha):
                self.linhas.append(self._linha)
            self._linha = None

    def handle_data(self, data):
        if self._celula is not None:
            self._celula.append(data)


def _linhas_do_arquivo(arquivo: Path) -> list[list[str]]:
    dados = Path(arquivo).read_bytes()
    if dados[:2] == b"PK":  # .xlsx é um zip
        from openpyxl import load_workbook

        planilha = load_workbook(io.BytesIO(dados), read_only=True, data_only=True).active
        return [["" if v is None else str(v).strip() for v in linha] for linha in planilha.iter_rows(values_only=True)]
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = dados.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if re.search(r"<\s*(table|tr)\b", texto[:5000], re.I):
        tabela = _TabelaHTML()
        tabela.feed(texto)
        return tabela.linhas
    primeira = texto.split("\n", 1)[0]
    separador = max([";", ",", "\t"], key=primeira.count)
    return [[c.strip() for c in linha] for linha in csv.reader(io.StringIO(texto), delimiter=separador)]


def _coluna(cabecalho: list[str], *chaves: str, evitar: tuple[str, ...] = ()) -> int | None:
    norm = [normalizar(c) for c in cabecalho]
    for chave in chaves:
        for i, c in enumerate(norm):
            if chave in c and not any(e in c for e in evitar):
                return i
    return None


def ler_exportacao(arquivo: Path) -> list[Matricula]:
    """Lê o arquivo exportado e devolve as matrículas (aluno, curso, progresso)."""
    linhas = _linhas_do_arquivo(arquivo)
    for n, cabecalho in enumerate(linhas[:10]):  # o cabeçalho pode vir depois de um título
        i_curso = _coluna(cabecalho, "curso", "treinamento", "produto")
        i_aluno = _coluna(cabecalho, "aluno", "nome completo", "usuario", "nome", "name",
                          evitar=("curso", "empresa", "filial", "sobrenome", "e-mail", "email"))
        if i_curso is not None and i_aluno is not None:
            break
    else:
        raise ValueError("A lista exportada não tem as colunas de aluno e curso. "
                         f"Cabeçalho encontrado: {linhas[0] if linhas else 'arquivo vazio'}")
    i_sobrenome = _coluna(cabecalho, "sobrenome", "last name")
    i_progresso = _coluna(cabecalho, "progresso", "progress", "conclus", "%")
    i_email = _coluna(cabecalho, "e-mail", "email")

    def celula(linha, i):
        return linha[i].strip() if i is not None and i < len(linha) and linha[i] else ""

    matriculas = []
    for linha in linhas[n + 1:]:
        aluno = " ".join(p for p in (celula(linha, i_aluno), celula(linha, i_sobrenome)) if p)
        if not aluno:
            continue
        texto = celula(linha, i_progresso)
        matriculas.append(Matricula(aluno, celula(linha, i_curso), ler_progresso(texto), texto, celula(linha, i_email)))
    return matriculas


# ---------- busca dos nomes ----------
def procurar(nomes: list[str], matriculas: list[Matricula]) -> list[Resultado]:
    """Para cada nome, as matrículas com o nome igual (sem acento/caixa).

    Se não houver nome igual, aceita quem tem todas as palavras do nome
    procurado (ex.: "Bruno Oliveira" acha "bruno matias goncalo de oliveira").
    """
    por_nome: dict[str, list[Matricula]] = {}
    for m in matriculas:
        por_nome.setdefault(normalizar(m.aluno), []).append(m)
    resultados = []
    for nome in nomes:
        chave = normalizar(nome)
        achadas = por_nome.get(chave)
        if not achadas:
            palavras = chave.split()
            achadas = [m for aluno, lista in por_nome.items()
                       if len(palavras) > 1 and all(p in aluno.split() for p in palavras) for m in lista]
        resultados.append(Resultado(nome, achadas or []))
    return resultados


def nome_certificado(aluno: str, curso: str) -> str:
    """'Nome - Curso.pdf', sem caracteres que o Windows não aceita em nome de arquivo."""
    aluno = " ".join(p if p in PARTICULAS else p.capitalize() for p in aluno.lower().split())
    nome = f"{aluno} - {curso.strip()}"
    nome = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", nome)
    nome = re.sub(r"\s+", " ", nome).strip(" .")
    return nome[:150] + ".pdf"
